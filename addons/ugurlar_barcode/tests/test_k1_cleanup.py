"""K1: performans/kalite denetimi P3 temizliği (TODO 33 + 37, borç B1/B2/Ö6).

Raf export O(n), rota listesi sayaçları tek sorgu, çok-worker çift okutma kilidi, etiket
barkodunun Odoo içinde üretimi, sayım null barkod, raflamada raf zorunluluğu.
"""
import base64
import hashlib
import io
import json
import random
import re
from types import SimpleNamespace
from unittest.mock import patch

from odoo import sql_db
from odoo.tests import tagged
from odoo.tools import file_open

from odoo.addons.odoougurlar.controllers import shelf_export
from odoo.addons.ugurlar_barcode.controllers import (
    base_api, batch_api, counting_api, operations_api, packing_api)
from odoo.addons.ugurlar_barcode.tests.common import BarcodeCase


def _src(name):
    with file_open('ugurlar_barcode/controllers/' + name) as f:
        return f.read()


# ═══ KK1: raf export toplamları ════════════════════════════════
@tagged('post_install', '-at_install')
class TestK1ShelfExport(BarcodeCase):

    def test_subtree_totals_equals_bruteforce_startswith(self):
        """Yeni sözlük toplamı, eski O(n²) startswith toplamıyla her lokasyonda aynıdır."""
        parent = self._new_shelf('K1-ANA', 'K1-ANA-1')
        for i in range(3):
            child = self._new_shelf('K1-ALT-%d' % i, 'K1-ALT-B%d' % i, parent=parent)
            self._new_shelf('K1-TORUN-%d' % i, 'K1-TORUN-B%d' % i, parent=child)
        locs = self.env['stock.location'].search([('usage', '=', 'internal')])
        paths = {loc.id: loc.parent_path or '' for loc in locs}
        rnd = random.Random(7)
        direct = {lid: rnd.choice((0, 1, 2.5, 4)) for lid in paths if rnd.random() < 0.7}

        totals = shelf_export._subtree_totals(paths, direct)

        self.assertGreaterEqual(len(paths), 7)
        for lid, pp in paths.items():
            expected = sum(direct.get(oid, 0) for oid, op in paths.items() if op.startswith(pp))
            self.assertAlmostEqual(totals.get(pp, 0.0), expected, msg='loc %s' % lid)

    def test_excel_total_quantity_includes_children(self):
        """Uçtan uca: Excel'deki Total Quantity alt raflar dahil, kendi dalıyla sınırlı."""
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        prod = self._new_product('K1 Export Ürün', 'K1-EXP-1')
        parent = self._new_shelf('K1-EXP-ANA', 'K1-EXP-A')
        c1 = self._new_shelf('K1-EXP-C1', 'K1-EXP-C1B', parent=parent)
        c2 = self._new_shelf('K1-EXP-C2', 'K1-EXP-C2B', parent=parent)
        other = self._new_shelf('K1-EXP-DIGER', 'K1-EXP-D')
        self._add_stock(prod, parent, 2)
        self._add_stock(prod, c1, 5)
        self._add_stock(prod, c2, 7)
        self._add_stock(prod, other, 100)
        self.env.flush_all()

        fake = SimpleNamespace(env=self.env, make_response=lambda data, headers=None: data)
        domain = json.dumps([['id', 'in', [parent.id, c1.id, c2.id, other.id]]])
        with patch.object(shelf_export, 'request', fake):
            data = shelf_export.ShelfExportController()._generate_shelf_excel(
                openpyxl, Font, PatternFill, Alignment, Border, Side, domain_str=domain)
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        total_by_id = {row[2]: row[5] for row in ws.iter_rows(min_row=2, values_only=True)}

        self.assertEqual(total_by_id[parent.id], 14)
        self.assertEqual(total_by_id[c1.id], 5)
        self.assertEqual(total_by_id[c2.id], 7)
        self.assertEqual(total_by_id[other.id], 100)


# ═══ KK2: rota listesi sayaçları ═══════════════════════════════
@tagged('post_install', '-at_install')
class TestK1BatchListCounters(BarcodeCase):

    def setUp(self):
        super().setUp()
        self._patch_request(batch_api)
        self.ctl = batch_api.BatchApiController()

    def _mk(self, n, state='draft'):
        Batch = self.env['stock.picking.batch']
        for i in range(n):
            batch = Batch.create({
                'picking_type_id': self.out_type.id,
                'time_window': '09:00-12:00',
                # yarısında source_info yok: depo adı picking_type'tan gelir
                'source_info': 'K1 DEPO - 1 siparis' if i % 2 else False,
            })
            if state != 'draft':
                # state hesaplanan+stored alan (picking'siz rota draft'a döner): doğrudan SQL
                self.env.flush_all()
                self.env.cr.execute(
                    'UPDATE stock_picking_batch SET state = %s WHERE id = %s', (state, batch.id))
                batch.invalidate_recordset(['state'])

    def _queries(self):
        self.env.flush_all()
        self.env.invalidate_all()
        seen = []
        real = self.env.cr.execute

        def spy(query, params=None, *a, **kw):
            seen.append(query)
            return real(query, params, *a, **kw)

        with patch.object(self.env.cr, 'execute', spy):
            res = self.ctl.batch_list()
        return res, len(seen)

    def test_state_counts_correct(self):
        base, _n = self._queries()
        self._mk(2)
        self._mk(1, 'in_progress')
        self._mk(3, 'done')
        res, _n = self._queries()
        for st, add in (('draft', 2), ('in_progress', 1), ('done', 3)):
            self.assertEqual(res['state_counts'][st], base['state_counts'][st] + add, st)

    def test_state_counts_one_grouped_query_not_per_state(self):
        """Durum başına search_count yok."""
        Batch = self.env['stock.picking.batch'].__class__
        with patch.object(Batch, 'search_count', side_effect=AssertionError('search_count')):
            res = self.ctl.batch_list()
        self.assertEqual(set(res['state_counts']), {'draft', 'in_progress', 'done'})

    def test_query_count_independent_of_batch_count(self):
        self.ctl.batch_list()   # ısınma: registry/ormcache ilk çağrıda ek sorgu çalıştırır
        self._mk(2)
        _res, few = self._queries()
        self._mk(8)
        res, many = self._queries()
        self.assertGreaterEqual(len(res['batches']), 10)
        self.assertEqual(few, many, 'sorgu sayısı rota sayısıyla artmamalı')


# ═══ KK3: çift okutma (çok worker) ═════════════════════════════
@tagged('post_install', '-at_install')
class TestK1DuplicateGuard(BarcodeCase):

    def setUp(self):
        super().setUp()
        # BarcodeCase pencereyi 0'lar; burada gerçek pencere gerekir
        window = patch.object(base_api, '_DEBOUNCE_SECONDS', 0.5)
        window.start()
        self.addCleanup(window.stop)
        saved = dict(base_api._last_requests)
        base_api._last_requests.clear()
        self.addCleanup(lambda: (base_api._last_requests.clear(),
                                 base_api._last_requests.update(saved)))
        self._patch_request()
        self.ctl = base_api.BarcodeApiBase()

    def _key(self, endpoint, *args):
        digest = hashlib.md5(str(args).encode()).hexdigest()
        return 'ugurlar_scan:%s:%s:%s' % (self.env.uid, endpoint, digest)

    def _other_worker_holds(self, key):
        """Ayrı bir bağlantı (= başka worker'ın transaction'ı) kilidi tutar."""
        cr = sql_db.db_connect(self.env.cr.dbname).cursor()
        self.addCleanup(cr.close)
        self.addCleanup(cr.rollback)
        cr.execute('SELECT pg_try_advisory_xact_lock(hashtextextended(%s, 0))', (key,))
        self.assertTrue(cr.fetchone()[0], 'ön koşul: diğer bağlantı kilidi aldı')

    def test_single_worker_behaviour_unchanged(self):
        self.assertFalse(self.ctl._check_duplicate_request('k1_scan', 1, 'K1-BC-A'))
        self.assertTrue(self.ctl._check_duplicate_request('k1_scan', 1, 'K1-BC-A'),
                        '500ms içinde aynı istek çift sayılır')
        self.assertFalse(self.ctl._check_duplicate_request('k1_scan', 1, 'K1-BC-B'),
                         'farklı barkod çift değil')

    def test_other_worker_in_flight_request_is_duplicate(self):
        self._other_worker_holds(self._key('k1_scan', 2, 'K1-BC-X'))
        self.assertTrue(self.ctl._check_duplicate_request('k1_scan', 2, 'K1-BC-X'))

    def test_other_worker_lock_does_not_block_different_barcode(self):
        self._other_worker_holds(self._key('k1_scan', 2, 'K1-BC-X'))
        self.assertFalse(self.ctl._check_duplicate_request('k1_scan', 2, 'K1-BC-Y'))

    def test_packing_and_collect_use_the_guard(self):
        self.assertIn("_check_duplicate_request('packing_scan'", _src('packing_api.py'))
        self.assertIn("_check_duplicate_request('batch_collect_scan'", _src('batch_api.py'))


# ═══ KK4: etiket barkodu Odoo içinde ═══════════════════════════
@tagged('post_install', '-at_install')
class TestK1LabelBarcode(BarcodeCase):

    def setUp(self):
        super().setUp()
        self._patch_request(packing_api)
        self.ctl = packing_api.PackingApiController()
        # HTTP'ye çıkılırsa test patlasın
        for name in ('urlopen', ):
            p = patch('urllib.request.' + name, side_effect=AssertionError('HTTP isteği'))
            p.start()
            self.addCleanup(p.stop)

    def _png(self, html):
        marker = 'data:image/png;base64,'
        start = html.index(marker) + len(marker)
        return base64.b64decode(html[start:html.index('"', start)])

    def test_template_label_barcode_is_data_uri_same_image(self):
        data = {'cargo_tracking': '1234567890'}
        elements = [{'type': 'cargo_barcode', 'x': 1, 'y': 1, 'width': 50, 'height': 15}]
        html = self.ctl._build_cargo_label_html(100, 150, elements, data)
        self.assertNotIn('/report/barcode', html)
        self.assertNotIn('http', html)
        expected = self.env['ir.actions.report'].barcode(
            'Code128', '1234567890', width=600, height=180, humanreadable=1)
        self.assertEqual(self._png(html), expected)

    def test_template_label_qr_is_data_uri(self):
        data = {'cargo_tracking': 'QR-123'}
        elements = [{'type': 'cargo_qr_code', 'x': 1, 'y': 1, 'width': 20, 'height': 20}]
        html = self.ctl._build_cargo_label_html(100, 150, elements, data)
        self.assertNotIn('/report/barcode', html)
        self.assertTrue(self._png(html).startswith(b'\x89PNG'))

    def test_default_label_barcode_is_data_uri_same_image(self):
        html = self.ctl._build_default_cargo_html(
            {'cargo_tracking': '9876543210', 'items': []})
        self.assertNotIn('/report/barcode', html)
        expected = self.env['ir.actions.report'].barcode(
            'Code128', '9876543210', width=840, height=200, humanreadable=1)
        self.assertEqual(self._png(html), expected)

    def test_no_tracking_no_image(self):
        elements = [{'type': 'cargo_barcode', 'x': 1, 'y': 1, 'width': 50, 'height': 15}]
        html = self.ctl._build_cargo_label_html(100, 150, elements, {'cargo_tracking': ''})
        self.assertNotIn('<img', html)

    def test_unencodable_barcode_skipped_not_raised(self):
        elements = [{'type': 'cargo_barcode', 'x': 1, 'y': 1, 'width': 50, 'height': 15}]
        with patch.object(type(self.env['ir.actions.report']), 'barcode',
                          side_effect=ValueError('x')):
            html = self.ctl._build_cargo_label_html(
                100, 150, elements, {'cargo_tracking': 'ÇÖ'})
        self.assertNotIn('<img', html)

    def test_label_code_does_not_build_report_barcode_urls(self):
        self.assertNotIn('_get_base_url', _src('packing_api.py'))
        self.assertNotIn('{base_url}', _src('packing_api.py'))


# ═══ KK5: paketleme log seviyesi ═══════════════════════════════
@tagged('post_install', '-at_install')
class TestK1PackingLogLevel(BarcodeCase):

    def test_packing_scan_detail_log_is_debug(self):
        src = _src('packing_api.py')
        self.assertNotIn('DEBUG LOG', src)
        match = re.search(r'_logger\.(\w+)\(\s*"PACKING SCAN %s', src)
        self.assertTrue(match, 'okutma ayrıntı logu bulunamadı')
        self.assertEqual(match.group(1), 'debug')


# ═══ KK6: sayımda null barkod ══════════════════════════════════
@tagged('post_install', '-at_install')
class TestK1CountNullBarcode(BarcodeCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.prod = cls._new_product('K1 Sayım Ürünü', 'K1-CN-1')
        cls.shelf = cls._new_shelf('K1-CN-RAF', 'K1-CN-RAF-1')

    def setUp(self):
        super().setUp()
        self._patch_request(counting_api)
        self.ctl = counting_api.CountingApiController()

    def test_null_barcode_item_is_not_found_others_still_counted(self):
        res = self.ctl.count_save(shelf_barcode='K1-CN-RAF-1', items=[
            {'barcode': None, 'quantity': 3},
            {'quantity': 2},
            {'barcode': self.prod.barcode, 'quantity': 4},
        ])
        self.assertTrue(res.get('success'), res)
        statuses = [r['status'] for r in res['results']]
        self.assertEqual(statuses[:2], ['not_found', 'not_found'])
        self.assertNotEqual(statuses[2], 'not_found')
        self.assertEqual(self._qty(self.prod, self.shelf), 4)


# ═══ KK7 / KK8: toplu raflama ══════════════════════════════════
@tagged('post_install', '-at_install')
class TestK1BulkPutaway(BarcodeCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.prod = cls._new_product('K1 Raflama Ürünü', 'K1-PW-1')
        cls.shelf = cls._new_shelf('K1-PW-RAF', 'K1-PW-RAF-1')

    def setUp(self):
        super().setUp()
        self._patch_request(operations_api)
        self.ctl = operations_api.OperationsApiController()

    def test_putaway_without_shelf_is_error(self):
        moves_before = self.env['stock.move'].search_count([('product_id', '=', self.prod.id)])
        for shelf in ('', None, '   '):
            res = self.ctl.bulk_scan(
                items=[{'barcode': self.prod.barcode, 'quantity': 2}],
                operation='putaway', shelf_barcode=shelf)
            self.assertIn('Raf gerekli', res.get('error', ''), repr(shelf))
            self.assertNotIn('results', res)
        self.assertEqual(
            self.env['stock.move'].search_count([('product_id', '=', self.prod.id)]), moves_before)

    def test_putaway_with_shelf_still_works(self):
        res = self.ctl.bulk_scan(
            items=[{'barcode': self.prod.barcode, 'quantity': 2}],
            operation='putaway', shelf_barcode='K1-PW-RAF-1')
        self.assertFalse(res.get('error'), res)
        self.assertEqual(res['success_count'], 1)
        self.assertEqual(self._qty(self.prod, self.shelf), 2)

    def test_info_without_shelf_unaffected(self):
        res = self.ctl.bulk_scan(items=[{'barcode': self.prod.barcode}], operation='info')
        self.assertFalse(res.get('error'), res)
        self.assertEqual(res['total'], 1)

    def test_dead_positive_filter_removed(self):
        src = _src('operations_api.py')
        self.assertNotIn('e[3] > 0', src)
        self.assertNotIn('qty <= 0', src)
